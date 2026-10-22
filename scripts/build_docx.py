#!/usr/bin/env python3
"""Build a styled Word document from the ADVP requirements Markdown.

Usage:
    python scripts/build_docx.py
    python scripts/build_docx.py --input docs/需求说明书.md --output docs/需求说明书.docx

Requires python-docx (already used by the project tooling).
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "docs" / "需求说明书.md"
DEFAULT_OUTPUT = ROOT / "docs" / "需求说明书.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
TEAL = RGBColor(0x1F, 0x6F, 0x6B)
GRAY = RGBColor(0x5A, 0x63, 0x70)
LIGHT = RGBColor(0x7A, 0x84, 0x92)
BODY_FONT = "Microsoft YaHei"
LATIN_FONT = "Segoe UI"
MONO_FONT = "Cascadia Code"


def set_run_font(run, name: str = BODY_FONT, size: float = 10.5,
                 bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), BODY_FONT if name not in (MONO_FONT,) else BODY_FONT)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top: int = 40, bottom: int = 40,
                     left: int = 70, right: int = 70) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_field(paragraph, instruction: str, size: float = 8.0,
              color: RGBColor | None = None, placeholder: str = "1") -> None:
    run = paragraph.add_run()
    set_run_font(run, size=size, color=color)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def add_page_number_footer(section, version: str) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第 ")
    set_run_font(run, size=8, color=LIGHT)
    add_field(p, "PAGE")
    run = p.add_run(" 页 / 共 ")
    set_run_font(run, size=8, color=LIGHT)
    add_field(p, "NUMPAGES")
    run = p.add_run(f" 页    ADVP-SRS-001 · v{version}")
    set_run_font(run, size=8, color=LIGHT)


def add_header(section, version: str) -> None:
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    run = p.add_run("ADVP 需求说明书")
    set_run_font(run, size=8.5, color=NAVY, bold=True)
    run = p.add_run("\t\tADVP-SRS-001 · v" + version)
    set_run_font(run, size=8.5, color=LIGHT)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.0))


def add_bookmark_heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    add_inline(p, text, base_size=16 if level == 1 else 13 if level == 2 else 11.5,
               base_bold=True, base_color=NAVY if level <= 2 else TEAL)


def add_inline(paragraph, text: str, base_size: float = 10.5,
               base_bold: bool = False, base_color: RGBColor | None = None) -> None:
    """Add text with **bold**, `code`, and [text](url) support."""
    token_re = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size, bold=base_bold, color=base_color)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=base_color)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=MONO_FONT, size=base_size - 0.8,
                         bold=base_bold, color=RGBColor(0x7A, 0x2E, 0x2E))
        else:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label, url = (m.group(1), m.group(2)) if m else (token, "")
            run = paragraph.add_run(label)
            set_run_font(run, size=base_size, bold=base_bold, color=RGBColor(0x1F, 0x5F, 0x9F))
            if url:
                run = paragraph.add_run(f" ({url})")
                set_run_font(run, size=base_size - 0.5, color=LIGHT)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, bold=base_bold, color=base_color)


def add_code_block(doc, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F4F7")
    p_pr.append(shd)
    run = p.add_run("\n".join(lines))
    set_run_font(run, name=MONO_FONT, size=8.5, color=RGBColor(0x24, 0x2A, 0x33))


def add_quote(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "35D0BA")
    borders.append(left)
    p_pr.append(borders)
    add_inline(p, text, base_size=9.5, base_color=GRAY)


def add_rule(doc) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D0D5DD")
    borders.append(bottom)
    p_pr.append(borders)


def split_row(line: str) -> list[str]:
    raw = line.strip().strip("|")
    return [cell.strip() for cell in raw.split("|")]


def is_separator(line: str) -> bool:
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def add_table(doc, rows: list[list[str]], header: list[str]) -> None:
    cols = max(len(header), max((len(r) for r in rows), default=0))
    if cols == 0:
        return
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_cell_margins(table)
    header_cells = table.rows[0].cells
    for idx in range(cols):
        text = header[idx] if idx < len(header) else ""
        cell = header_cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        add_inline(p, text, base_size=8.5, base_bold=True, base_color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "1F3864")
    repeat_header(table.rows[0])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx in range(cols):
            text = row[idx] if idx < len(row) else ""
            cell = cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, text, base_size=8.2, base_color=RGBColor(0x24, 0x2A, 0x33))
            if row_idx % 2 == 1:
                shade(cell, "F5F7FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_cover(doc, version: str, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("ADVP")
    set_run_font(run, size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    run = p.add_run("Automated Data Visualization Platform")
    set_run_font(run, size=12, color=TEAL)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    run = p.add_run(title)
    set_run_font(run, size=26, bold=True, color=RGBColor(0x11, 0x18, 0x27))

    p = doc.add_paragraph()
    run = p.add_run("Software Requirements Specification")
    set_run_font(run, size=12, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    run = p.add_run("文档编号  ADVP-SRS-001        版本  v" + version)
    set_run_font(run, size=10, color=GRAY)
    p = doc.add_paragraph()
    run = p.add_run("状态  待用户确认        日期  2026-10-22")
    set_run_font(run, size=10, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("本文件为 ADVP v1 需求基线草案，任何修改均递增 patch 版本并记录于 CHANGELOG。")
    set_run_font(run, size=9, color=LIGHT)

    doc.add_page_break()


def add_toc(doc) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("目录")
    set_run_font(run, size=16, bold=True, color=NAVY)
    p = doc.add_paragraph()
    run = p.add_run("提示：在 Word 中右键“更新域”可刷新目录与页码。")
    run.italic = True
    set_run_font(run, size=8.5, color=LIGHT)
    p = doc.add_paragraph()
    add_field(
        p,
        r'TOC \o "1-2" \h \z \u',
        size=10.5,
        color=GRAY,
        placeholder="（在 Word 中右键“更新域”生成目录）",
    )
    doc.add_page_break()


def configure_document(doc, version: str) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    add_header(section, version)
    add_page_number_footer(section, version)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)

    for lvl, size in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
        style = styles[f"Heading {lvl}"]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY if lvl <= 2 else TEAL
        style.paragraph_format.keep_with_next = True

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def parse_markdown(md: str) -> list[tuple[str, object]]:
    lines = md.splitlines()
    blocks: list[tuple[str, object]] = []
    i = 0
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []
    para_lines: list[str] = []

    def flush_para() -> None:
        nonlocal para_lines
        if para_lines:
            blocks.append(("p", " ".join(para_lines).strip()))
            para_lines = []

    def flush_table() -> None:
        nonlocal table_lines
        if not table_lines:
            return
        raw = [ln for ln in table_lines if ln.strip()]
        table_lines = []
        if len(raw) < 2:
            blocks.append(("p", " ".join(raw)))
            return
        header = split_row(raw[0])
        body = [split_row(ln) for ln in raw[2:]]
        blocks.append(("table", (header, body)))

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                blocks.append(("code", code_lines))
                code_lines = []
                in_code = False
            else:
                flush_para()
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_para()
            table_lines.append(stripped)
            i += 1
            continue
        if table_lines:
            flush_table()

        if not stripped:
            flush_para()
            i += 1
            continue

        if stripped.startswith("#"):
            flush_para()
            m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
            if m:
                blocks.append(("h", (len(m.group(1)), m.group(2))))
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            flush_para()
            blocks.append(("rule", None))
            i += 1
            continue

        if stripped.startswith(">"):
            flush_para()
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(("quote", " ".join(quote_lines)))
            continue

        m = re.match(r"^(\s*)([-*])\s+(.*)$", line)
        if m:
            flush_para()
            blocks.append(("bullet", (len(m.group(1)) // 2, m.group(3))))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            flush_para()
            blocks.append(("number", (len(m.group(1)) // 2, m.group(3))))
            i += 1
            continue

        para_lines.append(stripped)
        i += 1

    if in_code and code_lines:
        blocks.append(("code", code_lines))
    flush_table()
    flush_para()
    return blocks


def build(input_path: Path, output_path: Path) -> None:
    md = input_path.read_text(encoding="utf-8")
    version = (ROOT / "VERSION").read_text(encoding="utf-8").strip() if (ROOT / "VERSION").exists() else "0.1.0"
    blocks = parse_markdown(md)

    doc = Document()
    configure_document(doc, version)
    add_cover(doc, version, "需求说明书")
    add_toc(doc)

    skip_first_h1 = True
    for kind, payload in blocks:
        if kind == "h":
            level, text = payload  # type: ignore[misc]
            if level == 1 and skip_first_h1:
                skip_first_h1 = False
                continue
            add_bookmark_heading(doc, text, level)
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.0)
            add_inline(p, payload)  # type: ignore[arg-type]
        elif kind == "table":
            header, body = payload  # type: ignore[misc]
            add_table(doc, body, header)
        elif kind == "code":
            add_code_block(doc, payload)  # type: ignore[arg-type]
        elif kind == "quote":
            add_quote(doc, payload)  # type: ignore[arg-type]
        elif kind == "rule":
            add_rule(doc)
        elif kind in ("bullet", "number"):
            indent, text = payload  # type: ignore[misc]
            style = "List Bullet" if kind == "bullet" else "List Number"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.left_indent = Cm(0.6 + indent * 0.5)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    # validation
    check = Document(str(output_path))
    paragraphs = len(check.paragraphs)
    tables = len(check.tables)
    size = output_path.stat().st_size
    if paragraphs < 100:
        raise SystemExit(f"validation failed: only {paragraphs} paragraphs")
    if tables < 10:
        raise SystemExit(f"validation failed: only {tables} tables")
    if size < 50_000:
        raise SystemExit(f"validation failed: docx too small ({size} bytes)")
    print(f"OK: wrote {output_path}")
    print(f"  paragraphs: {paragraphs}")
    print(f"  tables: {tables}")
    print(f"  size: {size} bytes")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if not args.input.exists():
        print(f"FAIL: missing input {args.input}")
        return 1
    build(args.input, args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())

