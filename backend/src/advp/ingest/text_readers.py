"""Read local document text into heading-aware blocks."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import docx
import pdfplumber
from charset_normalizer import from_bytes
from lxml import html as lxml_html  # type: ignore[import-untyped]

from advp.ingest.files import (
    IngestError,
    InvalidSourceError,
    UnsupportedFormatError,
)

TEXT_SUFFIXES = {".txt", ".docx", ".pdf", ".md", ".rtf", ".html", ".htm"}
BULLET_RE = re.compile(r"^\s*(?:[·•●▪◦*\-]|\d+[.、)])\s*")
TERMINAL_PUNCTUATION = "。！？；;.!?"
HEADING_MAX_CHARS = 100


class TextDocumentEmptyError(InvalidSourceError):
    code = "document_text_empty"


class TextExtractionError(InvalidSourceError):
    code = "text_extraction_failed"


@dataclass(slots=True)
class TextLine:
    text: str
    line_no: int
    page: int | None = None


@dataclass(slots=True)
class TextBlock:
    heading: str | None
    heading_path: list[str]
    lines: list[TextLine]


@dataclass(slots=True)
class TextDocument:
    path: Path
    format: str
    blocks: list[TextBlock]
    char_count: int
    truncated: bool


def read_text_document(path_value: str | Path, *, max_chars: int) -> TextDocument:
    """Read a supported document into normalized heading blocks."""
    path = Path(path_value).expanduser().resolve()
    if not path.exists() or not path.is_file():
        raise InvalidSourceError(f"document was not found: {path.name}")
    suffix = path.suffix.lower()
    if suffix not in TEXT_SUFFIXES:
        raise UnsupportedFormatError(
            f"unsupported document format: {suffix or '<none>'}",
            hint="supported formats: TXT, DOCX, PDF, Markdown, RTF, HTML",
        )
    if max_chars < 1:
        raise TextDocumentEmptyError("text character limit must be positive")

    if suffix == ".docx":
        return _read_docx(path, max_chars=max_chars)
    if suffix == ".txt":
        return _read_txt(path, max_chars=max_chars)
    if suffix == ".pdf":
        return _read_pdf(path, max_chars=max_chars)
    if suffix == ".md":
        return _read_markdown(path, max_chars=max_chars)
    if suffix in {".html", ".htm"}:
        return _read_html(path, max_chars=max_chars)
    if suffix == ".rtf":
        return _read_rtf(path, max_chars=max_chars)
    raise UnsupportedFormatError(
        f"text reader is not available for {suffix}",
        hint="supported formats: TXT, DOCX, PDF, Markdown, RTF, HTML",
    )


def _read_txt(path: Path, *, max_chars: int) -> TextDocument:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise TextExtractionError("text file looks binary")
    text = _decode_text(raw)
    return _document_from_lines(
        path,
        format_name="txt",
        lines=text.splitlines(),
        max_chars=max_chars,
        styles=None,
        pages=None,
    )


def _read_docx(path: Path, *, max_chars: int) -> TextDocument:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise TextExtractionError(f"cannot read DOCX text: {exc}") from exc
    lines = [paragraph.text for paragraph in document.paragraphs]
    styles = [
        paragraph.style.name if paragraph.style is not None else ""
        for paragraph in document.paragraphs
    ]
    return _document_from_lines(
        path,
        format_name="docx",
        lines=lines,
        max_chars=max_chars,
        styles=styles,
        pages=None,
    )


def _read_pdf(path: Path, *, max_chars: int) -> TextDocument:
    lines: list[str] = []
    pages: list[int] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                for line in text.splitlines():
                    lines.append(line)
                    pages.append(page_number)
    except Exception as exc:
        raise TextExtractionError(f"cannot read PDF text: {exc}") from exc
    if not any(line.strip() for line in lines):
        raise TextDocumentEmptyError(
            "PDF contains no text layer",
            hint="scanned PDFs and images require OCR, which is not supported",
        )
    return _document_from_lines(
        path,
        format_name="pdf",
        lines=lines,
        max_chars=max_chars,
        styles=None,
        pages=pages,
    )


def _read_markdown(path: Path, *, max_chars: int) -> TextDocument:
    raw = _decode_text(path.read_bytes())
    lines: list[str] = []
    styles: list[str] = []
    in_code = False
    for line in raw.splitlines():
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            lines.append("")
            styles.append("")
            continue
        if in_code:
            lines.append("")
            styles.append("")
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading is not None:
            lines.append(heading.group(2))
            styles.append(f"Heading {len(heading.group(1))}")
            continue
        lines.append(line)
        styles.append("")
    return _document_from_lines(
        path,
        format_name="md",
        lines=lines,
        max_chars=max_chars,
        styles=styles,
        pages=None,
    )


def _read_html(path: Path, *, max_chars: int) -> TextDocument:
    raw = _decode_text(path.read_bytes())
    try:
        document = lxml_html.fromstring(raw)
    except Exception as exc:
        raise TextExtractionError(f"cannot parse HTML: {exc}") from exc
    for unwanted in document.xpath("//script | //style | //noscript | //template"):
        unwanted.drop_tree()
    body = document.body if document.body is not None else document
    lines: list[str] = []
    styles: list[str] = []
    for element in body.iter():
        tag = str(element.tag).lower() if isinstance(element.tag, str) else ""
        if tag not in {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li"}:
            continue
        text = " ".join(element.text_content().split())
        if not text:
            continue
        lines.append(text)
        styles.append(f"Heading {tag[1]}" if tag.startswith("h") else "")
    if not any(line.strip() for line in lines):
        text = " ".join(body.text_content().split())
        if text:
            lines = [text]
            styles = [""]
    return _document_from_lines(
        path,
        format_name="html",
        lines=lines,
        max_chars=max_chars,
        styles=styles,
        pages=None,
    )


def _read_rtf(path: Path, *, max_chars: int) -> TextDocument:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as exc:
        raise TextExtractionError(
            "RTF support is not installed",
            hint="install the striprtf dependency",
        ) from exc
    raw = path.read_bytes()
    try:
        decoded = raw.decode("utf-8")
    except UnicodeDecodeError:
        decoded = raw.decode("latin-1")
    text = rtf_to_text(decoded, errors="ignore")  # type: ignore[no-untyped-call]
    return _document_from_lines(
        path,
        format_name="rtf",
        lines=text.splitlines(),
        max_chars=max_chars,
        styles=None,
        pages=None,
    )


def _decode_text(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig")
    best = from_bytes(raw).best()
    if best is not None:
        return str(best)
    for encoding in ("utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise TextExtractionError("cannot decode text file")


def _document_from_lines(
    path: Path,
    *,
    format_name: str,
    lines: list[str],
    max_chars: int,
    styles: list[str] | None,
    pages: list[int] | None,
) -> TextDocument:
    normalized = [_normalize_line(line) for line in lines]
    full_text = "\n".join(normalized)
    truncated = len(full_text) > max_chars
    if truncated:
        full_text = full_text[:max_chars]
        normalized = full_text.split("\n")
        if pages is not None:
            pages = pages[: len(normalized)]
    char_count = len(full_text)
    if not full_text.strip():
        raise TextDocumentEmptyError("document contains no extractable text")

    blocks: list[TextBlock] = []
    current: TextBlock | None = None
    for index, line in enumerate(normalized):
        text = line.strip()
        if not text:
            continue
        style = styles[index] if styles and index < len(styles) else ""
        if _is_heading(text, style=style):
            heading = _heading_label(text)
            if not heading:
                continue
            current = TextBlock(heading=heading, heading_path=[heading], lines=[])
            blocks.append(current)
            continue
        if current is None:
            current = TextBlock(heading=None, heading_path=[], lines=[])
            blocks.append(current)
        page = pages[index] if pages is not None and index < len(pages) else None
        current.lines.append(TextLine(text=text, line_no=index + 1, page=page))

    blocks = [block for block in blocks if block.lines]
    if not blocks:
        raise TextDocumentEmptyError("document contains no extractable text")
    return TextDocument(
        path=path,
        format=format_name,
        blocks=blocks,
        char_count=char_count,
        truncated=truncated,
    )


def _normalize_line(line: str) -> str:
    return (
        line.replace("\u3000", " ")
        .replace("\xa0", " ")
        .replace("：", ":")
        .replace("，", ", ")
        .replace("；", "; ")
        .rstrip()
    )


def _is_heading(text: str, *, style: str) -> bool:
    style_lower = style.lower()
    if style_lower.startswith("heading") or "标题" in style:
        return True
    if BULLET_RE.match(text):
        return False
    outside = re.sub(r"[（(][^（）()]*[)）]", "", text).strip()
    if ":" in outside:
        return False
    if len(outside) > HEADING_MAX_CHARS:
        return False
    return not any(mark in outside for mark in TERMINAL_PUNCTUATION)


def _heading_label(text: str) -> str:
    label = re.sub(r"[（(].*?[)）]", "", text).strip()
    return label or text.strip()


__all__ = [
    "IngestError",
    "TextBlock",
    "TextDocument",
    "TextDocumentEmptyError",
    "TextExtractionError",
    "TextLine",
    "read_text_document",
]
